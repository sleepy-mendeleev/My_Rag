"""
文档加载器模块
支持多种文档格式的加载和预处理
"""

import os
import logging
import hashlib
from typing import List, Dict, Any, Optional
from pathlib import Path

# LangChain imports
from langchain_core.documents import Document
from langchain_community.document_loaders import (
    TextLoader,
    CSVLoader,
    DirectoryLoader
)

# 配置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocumentLoader:
    """文档加载器类"""
    
    # 支持的文件类型
    SUPPORTED_EXTENSIONS = {
        '.txt': 'text',
        '.csv': 'csv',
        '.pdf': 'pdf',
        '.docx': 'docx',
        '.xlsx': 'excel',
        '.xls': 'excel'
    }
    
    def __init__(self, encoding: str = 'utf-8'):
        """
        初始化文档加载器
        
        Args:
            encoding: 文件编码格式
        """
        self.encoding = encoding
    
    def get_file_type(self, file_path: str) -> str:
        """
        获取文件类型
        
        Args:
            file_path: 文件路径
            
        Returns:
            文件类型
        """
        extension = Path(file_path).suffix.lower()
        return self.SUPPORTED_EXTENSIONS.get(extension, 'unknown')
    
    def is_supported(self, file_path: str) -> bool:
        """
        检查文件是否支持
        
        Args:
            file_path: 文件路径
            
        Returns:
            是否支持
        """
        return self.get_file_type(file_path) != 'unknown'
    
    def load_text_file(self, file_path: str) -> List[Document]:
        """加载文本文件"""
        try:
            loader = TextLoader(file_path, encoding=self.encoding)
            documents = loader.load()
            
            # 添加文件信息到元数据
            for doc in documents:
                doc.metadata.update({
                    'source': file_path,
                    'file_type': 'text',
                    'file_name': Path(file_path).name
                })
            
            return documents
        except Exception as e:
            logger.error(f"加载文本文件失败 {file_path}: {e}")
            return []
    
    def load_csv_file(self, file_path: str, **kwargs) -> List[Document]:
        """读取 CSV，保留表头和二维数据行，交给表格切分器处理。"""
        try:
            import pandas as pd

            encoding = self.encoding
            if encoding.lower().replace("_", "-") in ("utf-8", "utf8"):
                encoding = "utf-8-sig"

            options = {
                "encoding": encoding,
                "dtype": str,
                "keep_default_na": False,
            }
            options.update(kwargs)
            frame = pd.read_csv(file_path, **options)

            if frame.empty:
                raise ValueError("CSV 没有数据记录")

            return [Document(
                page_content=frame.to_csv(index=False),
                metadata={
                    "source": file_path,
                    "file_type": "csv",
                    "file_name": Path(file_path).name,
                    "content_type": "table",
                    "table_data": {
                        "headers": frame.columns.tolist(),
                        "rows": frame.values.tolist(),
                    },
                },
            )]
        except Exception as e:
            logger.error(f"加载 CSV 失败 {file_path}: {e}")
            raise
    
    def load_pdf_file(self, file_path: str) -> List[Document]:
        """按原文顺序封装正文与表格，交给统一切分器。"""
        try:
            from .pdf_parser import extract_pdf_pages

            pages = extract_pdf_pages(file_path)
            documents = []

            for page in pages:

                base_metadata = {
                    "source": file_path,
                    "file_type": "pdf",
                    "file_name": Path(file_path).name,
                    "page_number": page["page_number"],
                    "extraction_method": page["extraction_method"],
                }
                buffer = []

                def flush_text():
                    if buffer:
                        documents.append(Document(
                            page_content="\n\n".join(buffer),
                            metadata={**base_metadata, "content_type": "text"},
                        ))
                        buffer.clear()

                for block in page["blocks"]:
                    if "table_data" in block:
                        flush_text()
                        table_metadata = {
                            **base_metadata,
                            "content_type": "table",
                            "table_number": block["table_number"],
                            "table_bbox": block["bbox"],
                            "table_data": block["table_data"],
                        }
                        for key in ("table_title", "table_note"):
                            if block.get(key):
                                table_metadata[key] = block[key]
                        documents.append(Document(
                            page_content=block["text"],
                            metadata=table_metadata,
                        ))
                    else:
                        buffer.append(block["text"])
                flush_text()

            return documents

        except Exception as e:
            logger.error(f"加载 PDF 失败 {file_path}: {e}")
            raise
    
    def load_docx_file(self, file_path: str) -> List[Document]:
        """按原文顺序提取 Word 正文和普通表格。"""
        try:
            import re
            from docx import Document as WordDocument
            from docx.text.paragraph import Paragraph
            from docx.table import Table

            base_metadata = {
                "source": file_path,
                "file_type": "docx",
                "file_name": Path(file_path).name,
            }
            documents = []

            for block in WordDocument(file_path).iter_inner_content():
                if isinstance(block, Paragraph):
                    text = block.text.strip()
                    if not text:
                        continue

                    style_name = block.style.name if block.style else ""
                    heading = re.fullmatch(r"Heading ([1-6])", style_name)
                    if heading:
                        text = "#" * int(heading.group(1)) + " " + text

                    documents.append(Document(
                        page_content=text + "\n\n",
                        metadata={**base_metadata, "content_type": "text"},
                    ))

                elif isinstance(block, Table):
                    if block._tbl.xpath(".//w:gridSpan | .//w:vMerge"):
                        raise ValueError("当前步骤暂不支持含合并单元格的 Word 表格")
                    if any(cell.tables for row in block.rows for cell in row.cells):
                        raise ValueError("当前步骤暂不支持嵌套 Word 表格")

                    rows = [
                        [cell.text.strip() for cell in row.cells]
                        for row in block.rows
                    ]
                    if len(rows) < 2 or not any(any(row) for row in rows[1:]):
                        continue
                    if len({len(row) for row in rows}) != 1:
                        raise ValueError("Word 表格各行列数不一致")

                    documents.append(Document(
                        page_content="\n".join("\t".join(row) for row in rows),
                        metadata={
                            **base_metadata,
                            "content_type": "table",
                            "table_data": {
                                "headers": rows[0],
                                "rows": rows[1:],
                            },
                        },
                    ))

            if not documents:
                raise ValueError("Word 没有可提取的正文或表格数据")
            return documents

        except Exception as e:
            logger.error(f"加载 Word 失败 {file_path}: {e}")
            raise
    
    def load_excel_file(self, file_path: str) -> List[Document]:
        """读取各工作表，保留表头和完整数据行。"""
        try:
            import pandas as pd

            extension = Path(file_path).suffix.lower()
            engine = "xlrd" if extension == ".xls" else "openpyxl"

            documents = []

            with pd.ExcelFile(file_path, engine=engine) as workbook:
                for sheet_name in workbook.sheet_names:
                    frame = pd.read_excel(
                        workbook,
                        sheet_name=sheet_name,
                        dtype=str,
                        keep_default_na=False,
                    )

                    # 空工作表或只有表头的工作表不生成片段。
                    if frame.empty:
                        continue

                    documents.append(Document(
                        page_content=frame.to_csv(index=False),
                        metadata={
                            "source": file_path,
                            "file_type": "excel",
                            "file_name": Path(file_path).name,
                            "sheet_name": sheet_name,
                            "content_type": "table",
                            "table_data": {
                                "headers": frame.columns.tolist(),
                                "rows": frame.values.tolist(),
                            },
                        },
                    ))

            if not documents:
                raise ValueError("Excel 没有可入库的数据记录")

            return documents

        except Exception as e:
            logger.error(f"加载 Excel 失败 {file_path}: {e}")
            raise
    
    def load_single_file(self, file_path: str, **kwargs) -> List[Document]:
        """
        加载单个文件
        
        Args:
            file_path: 文件路径
            **kwargs: 额外参数
            
        Returns:
            文档列表
        """
        if not os.path.exists(file_path):
            logger.error(f"文件不存在: {file_path}")
            return []
        
        file_type = self.get_file_type(file_path)
        
        if file_type == 'text':
            documents = self.load_text_file(file_path)
        elif file_type == 'csv':
            documents = self.load_csv_file(file_path, **kwargs)
        elif file_type == 'pdf':
            documents = self.load_pdf_file(file_path)
        elif file_type == 'docx':
            documents = self.load_docx_file(file_path)
        elif file_type == 'excel':
            documents = self.load_excel_file(file_path)
        else:
            raise ValueError(f'暂不支持文件类型：{Path(file_path).suffix}；旧版 .doc 请先转为 .docx')
        if not documents or not any(doc.page_content.strip() for doc in documents):
            raise ValueError('未提取到有效文字，请检查文件格式、编码或扫描件 OCR 需求')
        # Content identity is stable when a file is uploaded to another temporary path.
        digest = hashlib.sha256()
        with open(file_path, 'rb') as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(block)
        for doc in documents:
            doc.metadata.update({'document_id': digest.hexdigest(), 'parser_version': 'structure-v1',
                                 'extraction_method': doc.metadata.get('extraction_method', 'native')})
        return documents
    
    def load_directory(self, directory_path: str, 
                      glob_pattern: str = "**/*",
                      exclude_patterns: List[str] = None) -> List[Document]:
        """
        加载目录中的所有支持文件
        
        Args:
            directory_path: 目录路径
            glob_pattern: 文件匹配模式
            exclude_patterns: 排除的文件模式
            
        Returns:
            文档列表
        """
        if not os.path.exists(directory_path):
            logger.error(f"目录不存在: {directory_path}")
            return []
        
        all_documents = []
        exclude_patterns = exclude_patterns or []
        
        try:
            # 遍历目录中的所有文件
            for file_path in Path(directory_path).glob(glob_pattern):
                if file_path.is_file():
                    # 检查是否需要排除
                    should_exclude = any(
                        pattern in str(file_path) for pattern in exclude_patterns
                    )
                    
                    if should_exclude:
                        continue
                    
                    # 检查文件类型是否支持
                    if self.is_supported(str(file_path)):
                        logger.info(f"加载文件: {file_path}")
                        documents = self.load_single_file(str(file_path))
                        all_documents.extend(documents)
                    else:
                        logger.debug(f"跳过不支持的文件: {file_path}")
            
            logger.info(f"从目录 {directory_path} 加载了 {len(all_documents)} 个文档")
            return all_documents
            
        except Exception as e:
            logger.error(f"加载目录失败 {directory_path}: {e}")
            return []
    
    def load_academic_csv(self, csv_path: str,
                         text_columns: List[str] = None,
                         metadata_columns: List[str] = None) -> List[Document]:
        """
        加载通用CSV文件并构建文档
        """
        try:
            import pandas as pd
            df = pd.read_csv(csv_path, encoding=self.encoding)
            logger.info(f"读取CSV文件: {csv_path}, 共 {len(df)} 行数据")

            if text_columns is None or not text_columns:
                object_cols = [c for c in df.columns if df[c].dtype == 'object']
                text_columns = [c for c in object_cols if not str(c).startswith('Unnamed')]

            if metadata_columns is None:
                metadata_columns = [c for c in df.columns if c not in text_columns]

            documents = []
            for idx, row in df.iterrows():
                content_parts = []
                metadata = {
                    'source': csv_path,
                    'file_type': 'csv',
                    'row_index': idx
                }

                for col in metadata_columns:
                    val = row.get(col)
                    if pd.notna(val):
                        metadata[str(col)] = str(val)

                for col in text_columns:
                    val = row.get(col)
                    if pd.notna(val):
                        text = str(val).strip()
                        if text:
                            content_parts.append(f"{col}: {text}")

                if content_parts:
                    content = "\n\n".join(content_parts)
                    documents.append(Document(page_content=content, metadata=metadata))

            logger.info(f"从CSV构建了 {len(documents)} 个文档")
            return documents
        except Exception as e:
            logger.error(f"加载CSV失败: {e}")
            return []
    
    def get_file_info(self, file_path: str) -> Dict[str, Any]:
        """
        获取文件信息
        
        Args:
            file_path: 文件路径
            
        Returns:
            文件信息字典
        """
        if not os.path.exists(file_path):
            return {"error": "文件不存在"}
        
        file_stat = os.stat(file_path)
        path_obj = Path(file_path)
        
        return {
            "file_name": path_obj.name,
            "file_path": str(path_obj.absolute()),
            "file_size": file_stat.st_size,
            "file_type": self.get_file_type(file_path),
            "is_supported": self.is_supported(file_path),
            "extension": path_obj.suffix.lower(),
            "created_time": file_stat.st_ctime,
            "modified_time": file_stat.st_mtime
        }


def main():
    """测试函数"""
    loader = DocumentLoader()
    
    # 测试加载学术CSV数据
    csv_path = "../data.csv"
    if os.path.exists(csv_path):
        documents = loader.load_academic_csv(csv_path)
        print(f"加载了 {len(documents)} 个文档")
        
        if documents:
            print("\n第一个文档示例:")
            print(f"内容: {documents[0].page_content[:300]}...")
            print(f"元数据: {documents[0].metadata}")
    
    # 测试文件信息获取
    print(f"\n文件信息:")
    info = loader.get_file_info(csv_path)
    print(info)


if __name__ == "__main__":
    main()
